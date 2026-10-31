from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from core.packets import Packet, PacketStatus

SUPPORTED_FORMATS = {".docx", ".xlsx", ".pptx"}

LEGACY_FORMATS = {".doc", ".xls", ".ppt"}

def _normalise_value(value):
    """Convert Office values into JSON-serialisable values."""

    if isinstance(value, (datetime, date, time)):
        return value.isoformat()

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, bytes):
        return value.hex()

    if isinstance(value, (str, int, float, bool)) or value is None:
        return value

    return str(value)

def _file_metadata(path):
    """Return common filesystem metadata."""

    return {
        "name": path.name,
        "extension": path.suffix.lower(),
        "absolute_path": str(path.resolve()),
        "size_bytes": path.stat().st_size,
    }

def _extract_docx(path):
    """Extract content and metadata from a Word document."""

    document = Document(path)
    properties = document.core_properties

    metadata = {
        "title": properties.title,
        "subject": properties.subject,
        "author": properties.author,
        "keywords": properties.keywords,
        "category": properties.category,
        "comments": properties.comments,
        "created": _normalise_value(properties.created),
        "modified": _normalise_value(properties.modified),
        "last_modified_by": properties.last_modified_by,
        "revision": properties.revision,
        "identifier": properties.identifier,
        "language": properties.language,
    }

    paragraphs = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        hyperlinks = []

        for run in paragraph.runs:
            relationship_id = run._r.get(
                "{http://schemas.openxmlformats.org/"
                "officeDocument/2006/relationships}id"
            )

            if relationship_id:
                relationship = document.part.rels.get(relationship_id)

                if relationship:
                    hyperlinks.append({
                        "text": run.text,
                        "url": relationship.target_ref,
                    })

        paragraphs.append({
            "index": index,
            "text": paragraph.text,
            "style": paragraph.style.name if paragraph.style else None,
            "hyperlinks": hyperlinks,
        })

    tables = []

    for table_index, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            rows.append([
                cell.text
                for cell in row.cells
            ])

        tables.append({
            "index": table_index,
            "row_count": len(rows),
            "column_count": max(
                (len(row) for row in rows),
                default=0,
            ),
            "rows": rows,
        })

    return {
        "format": "docx",
        "metadata": metadata,
        "content": {
            "paragraphs": paragraphs,
            "tables": tables,
        },
        "summary": {
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "word_count": sum(
                len(paragraph["text"].split())
                for paragraph in paragraphs
            ),
            "character_count": sum(
                len(paragraph["text"])
                for paragraph in paragraphs
            ),
        },
    }

def _extract_xlsx(path):
    """Extract workbook metadata, worksheets and cell contents."""

    workbook = load_workbook(
        filename=path,
        read_only=False,
        data_only=False,
        keep_links=True,
    )

    try:
        properties = workbook.properties

        metadata = {
            "title": properties.title,
            "subject": properties.subject,
            "creator": properties.creator,
            "description": properties.description,
            "keywords": properties.keywords,
            "category": properties.category,
            "created": _normalise_value(properties.created),
            "modified": _normalise_value(properties.modified),
            "last_modified_by": properties.lastModifiedBy,
            "revision": properties.revision,
            "identifier": properties.identifier,
            "language": properties.language,
        }

        worksheets = []
        total_cells = 0
        total_formulas = 0

        for sheet in workbook.worksheets:
            cells = []
            formula_count = 0
            populated_count = 0

            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue

                    value = _normalise_value(cell.value)
                    is_formula = (
                        isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    )

                    if is_formula:
                        formula_count += 1

                    populated_count += 1

                    cells.append({
                        "coordinate": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                        "value": value,
                        "data_type": cell.data_type,
                        "is_formula": is_formula,
                        "number_format": cell.number_format,
                        "hyperlink": (
                            cell.hyperlink.target
                            if cell.hyperlink
                            else None
                        ),
                    })

            merged_ranges = [
                str(merged_range)
                for merged_range in sheet.merged_cells.ranges
            ]

            worksheets.append({
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "freeze_panes": (
                    str(sheet.freeze_panes)
                    if sheet.freeze_panes
                    else None
                ),
                "merged_ranges": merged_ranges,
                "cells": cells,
                "summary": {
                    "populated_cells": populated_count,
                    "formula_count": formula_count,
                },
            })

            total_cells += populated_count
            total_formulas += formula_count

        return {
            "format": "xlsx",
            "metadata": metadata,
            "content": {
                "worksheets": worksheets,
            },
            "summary": {
                "worksheet_count": len(worksheets),
                "populated_cells": total_cells,
                "formula_count": total_formulas,
            },
        }

    finally:
        workbook.close()

def _extract_pptx(path):
    """Extract presentation metadata, slides and their contents."""

    presentation = Presentation(path)
    properties = presentation.core_properties

    metadata = {
        "title": properties.title,
        "subject": properties.subject,
        "author": properties.author,
        "keywords": properties.keywords,
        "comments": properties.comments,
        "category": properties.category,
        "created": _normalise_value(properties.created),
        "modified": _normalise_value(properties.modified),
        "last_modified_by": properties.last_modified_by,
        "revision": properties.revision,
        "identifier": properties.identifier,
        "language": properties.language,
    }

    slides = []
    total_text_length = 0
    total_shapes = 0
    total_tables = 0

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        shapes = []
        slide_text = []

        for shape in slide.shapes:
            shape_data = {
                "name": shape.name,
                "type": str(shape.shape_type),
                "shape_id": shape.shape_id,
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height,
                "text": None,
                "table": None,
            }

            if shape.has_text_frame:
                text = shape.text
                shape_data["text"] = text

                if text.strip():
                    slide_text.append(text)

            if shape.has_table:
                table_rows = []

                for row in shape.table.rows:
                    table_rows.append([
                        cell.text
                        for cell in row.cells
                    ])

                shape_data["table"] = {
                    "row_count": len(table_rows),
                    "column_count": max(
                        (len(row) for row in table_rows),
                        default=0,
                    ),
                    "rows": table_rows,
                }

                total_tables += 1

            shapes.append(shape_data)

        notes = None

        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame

            if notes_frame:
                notes = notes_frame.text

        text_content = "\n".join(slide_text)
        total_text_length += len(text_content)
        total_shapes += len(shapes)

        slides.append({
            "slide_number": slide_number,
            "layout": slide.slide_layout.name,
            "text": text_content,
            "shapes": shapes,
            "shape_count": len(shapes),
            "notes": notes,
        })

    return {
        "format": "pptx",
        "metadata": metadata,
        "presentation": {
            "slide_width": presentation.slide_width,
            "slide_height": presentation.slide_height,
            "slide_count": len(slides),
        },
        "content": {
            "slides": slides,
        },
        "summary": {
            "slide_count": len(slides),
            "shape_count": total_shapes,
            "table_count": total_tables,
            "text_length": total_text_length,
        },
    }

def officeextract(packet: Packet) -> Packet:
    """
    Extract metadata and content from supported Office documents.

    Supported:
        DOCX, XLSX, PPTX

    Legacy formats:
        DOC, XLS, PPT (not supported)
    """

    result = packet.child(
        stage="extraction",
        artefact_type="office",
        status=PacketStatus.RUNNING,
    )

    try:
        if not isinstance(packet.input, str):
            raise TypeError(
                "Office input must be a filesystem path."
            )

        path = Path(packet.input.strip()).expanduser()

        if not path.exists():
            raise FileNotFoundError(
                f"Office document does not exist: {path}"
            )

        if not path.is_file():
            raise ValueError(
                "Input must reference a regular file."
            )

        extension = path.suffix.lower()

        if extension in LEGACY_FORMATS:
            result.status = PacketStatus.NOT_SUPPORTED
            result.add_error(
                code="OFFICE_FORMAT_NOT_SUPPORTED",
                message=(
                    f"Legacy Office format '{extension}' "
                    "is not currently supported."
                ),
                component="extractors.office",
            )
            return result

        if extension not in SUPPORTED_FORMATS:
            raise ValueError(
                f"Unsupported Office document extension: {extension}"
            )

        extractors = {
            ".docx": _extract_docx,
            ".xlsx": _extract_xlsx,
            ".pptx": _extract_pptx,
        }

        extracted = extractors[extension](path)

        result.data = {
            "file": _file_metadata(path),
            **extracted,
        }

        result.status = PacketStatus.SUCCESS

    except FileNotFoundError as exc:
        result.status = PacketStatus.FAILED
        result.add_error(
            code="OFFICE_FILE_NOT_FOUND",
            message=str(exc),
            component="extractors.office",
        )

    except PermissionError as exc:
        result.status = PacketStatus.FAILED
        result.add_error(
            code="OFFICE_ACCESS_DENIED",
            message=str(exc),
            component="extractors.office",
        )

    except (TypeError, ValueError) as exc:
        result.status = PacketStatus.FAILED
        result.add_error(
            code="INVALID_OFFICE_INPUT",
            message=str(exc),
            component="extractors.office",
        )

    except Exception as exc:
        result.status = PacketStatus.FAILED
        result.add_error(
            code="OFFICE_EXTRACTION_FAILED",
            message=str(exc),
            component="extractors.office",
        )

    return result